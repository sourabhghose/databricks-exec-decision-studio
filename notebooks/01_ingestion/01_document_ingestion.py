# Databricks notebook source
# MAGIC %md
# MAGIC # Executive Decision Studio — Document Ingestion
# MAGIC **Alinta Energy | Databricks Executive Decision Studio**
# MAGIC
# MAGIC Reads raw document files from the Unity Catalog volume at
# MAGIC `/Volumes/ausnet_process_intel_catalog/eds_raw/documents/tier{N}/`,
# MAGIC parses each format to recover content and metadata, and stages the
# MAGIC documents in `ausnet_process_intel_catalog.eds_raw.staged_documents`
# MAGIC for downstream chunking and embedding.
# MAGIC
# MAGIC **Supported formats:** `.txt` (YAML frontmatter), `.pdf`, `.pptx`, `.xlsx`
# MAGIC
# MAGIC **Source**: `/Volumes/ausnet_process_intel_catalog/eds_raw/documents/`
# MAGIC **Target**: `ausnet_process_intel_catalog.eds_raw.staged_documents`

# COMMAND ----------

# MAGIC %pip install pdfplumber python-pptx openpyxl --quiet

# COMMAND ----------

dbutils.library.restartPython()

# COMMAND ----------

import uuid
import re
import os
from datetime import datetime, date
from pyspark.sql import functions as F, Row
from pyspark.sql.types import (
    StructType, StructField, StringType, IntegerType,
    BooleanType, DateType, TimestampType
)

CATALOG    = "ausnet_process_intel_catalog"
VOLUME_ROOT = f"/Volumes/{CATALOG}/eds_raw/documents"
TIERS       = ["tier1", "tier2", "tier3", "tier4"]
EXTENSIONS  = {".txt", ".pdf", ".pptx", ".xlsx"}

spark.sql(f"USE CATALOG {CATALOG}")

print("=" * 65)
print("EDS — Document Ingestion from Volume (multi-format)")
print("=" * 65)
print(f"Source  : {VOLUME_ROOT}")
print(f"Formats : {', '.join(sorted(EXTENSIONS))}")
print(f"Started : {datetime.utcnow().isoformat()}")

# COMMAND ----------

# MAGIC %md
# MAGIC ## Step 1: Create / verify eds_raw.staged_documents table

# COMMAND ----------

spark.sql(f"""
CREATE TABLE IF NOT EXISTS {CATALOG}.eds_raw.staged_documents (
    ingest_id           STRING      NOT NULL COMMENT 'Unique ingestion event ID',
    doc_id              STRING      NOT NULL COMMENT 'Source document ID',
    title               STRING      COMMENT 'Document title',
    doc_type            STRING      COMMENT 'Document type',
    classification      STRING      COMMENT 'Security classification',
    access_tier_level   INT         COMMENT 'Access tier: 1 (Board) to 4 (All staff)',
    business_area       STRING      COMMENT 'Business area',
    effective_date      DATE        COMMENT 'Document effective date',
    author              STRING      COMMENT 'Document author',
    version             STRING      COMMENT 'Document version',
    content             STRING      COMMENT 'Full document text (frontmatter stripped)',
    source_file         STRING      COMMENT 'Volume path of the source file',
    source_system       STRING      COMMENT 'Source system',
    word_count          INT         COMMENT 'Approximate word count',
    char_count          INT         COMMENT 'Total character count',
    ingest_timestamp    TIMESTAMP   COMMENT 'Ingestion pipeline timestamp',
    ingest_pipeline     STRING      COMMENT 'Pipeline that produced this record',
    ingest_run_id       STRING      COMMENT 'Run ID for this ingestion batch',
    is_synthetic        BOOLEAN     COMMENT 'True for synthetic content',
    ingestion_status    STRING      COMMENT 'staged | processed | failed'
)
USING DELTA
COMMENT 'Raw staged documents ingested from Unity Catalog volume — all formats'
TBLPROPERTIES (
    'delta.enableChangeDataFeed' = 'true',
    'delta.autoOptimize.optimizeWrite' = 'true'
)
""")

print(f"[OK] {CATALOG}.eds_raw.staged_documents ready")

# Backwards-compatible: add source_file column if it doesn't exist yet
try:
    spark.sql(f"""
        ALTER TABLE {CATALOG}.eds_raw.staged_documents
        ADD COLUMN source_file STRING COMMENT 'Volume path of the source file'
    """)
    print("[OK] Added source_file column (migration)")
except Exception:
    pass  # column already exists

# COMMAND ----------

# MAGIC %md
# MAGIC ## Step 2: Load metadata lookup from eds_synthetic.documents
# MAGIC
# MAGIC PDF, PPTX, and XLSX files do not carry embedded YAML frontmatter.
# MAGIC Metadata is recovered from `eds_synthetic.documents` keyed on `doc_id`
# MAGIC (extracted from the filename stem, e.g. `DOC-016.pdf` → `DOC-016`).

# COMMAND ----------

def load_metadata_lookup():
    """
    Return a dict keyed on doc_id with metadata rows from eds_synthetic.documents.
    Falls back to empty dict if the table does not exist.
    """
    try:
        rows = spark.sql(f"""
            SELECT doc_id, title, doc_type, classification,
                   access_tier_level, business_area, effective_date,
                   author, version, source_system, is_synthetic
            FROM {CATALOG}.eds_synthetic.documents
        """).collect()
        return {r["doc_id"]: r for r in rows}
    except Exception as e:
        print(f"[WARN] Could not load eds_synthetic.documents: {e}")
        return {}

METADATA_LOOKUP = load_metadata_lookup()
print(f"[OK] Loaded {len(METADATA_LOOKUP)} document metadata records from eds_synthetic.documents")

# COMMAND ----------

# MAGIC %md
# MAGIC ## Step 3: File discovery — scan all tier folders

# COMMAND ----------

def list_volume_files():
    """
    Return list of (tier, posix_path, extension) tuples for all
    supported files across all tier subdirectories.
    Uses dbutils.fs.ls for directory listing; paths converted to POSIX
    (removing dbfs: prefix where applicable) for binary open() calls.
    """
    files = []
    for tier in TIERS:
        tier_path = f"{VOLUME_ROOT}/{tier}"
        try:
            entries = dbutils.fs.ls(tier_path)
        except Exception as e:
            print(f"[WARN] Cannot list {tier_path}: {e}")
            continue
        for entry in entries:
            name = entry.name.rstrip("/")
            ext = os.path.splitext(name)[1].lower()
            if ext in EXTENSIONS:
                # Unity Catalog volumes are accessible as POSIX paths directly
                posix_path = f"{VOLUME_ROOT}/{tier}/{name}"
                files.append((tier, posix_path, ext))
    return files

volume_files = list_volume_files()
print(f"\nFound {len(volume_files)} files across tiers {TIERS}")
for tier, path, ext in volume_files:
    print(f"  [{tier}] {os.path.basename(path):35s}  ({ext})")

# COMMAND ----------

# MAGIC %md
# MAGIC ## Step 4: Parser functions

# COMMAND ----------

# ---- Shared helpers ----

FRONTMATTER_RE = re.compile(r"^---\n(.*?)\n---\n", re.DOTALL)

def safe_int(val, default=4):
    try:
        return int(val)
    except (TypeError, ValueError):
        return default

def safe_date(val):
    if val is None:
        return None
    if isinstance(val, date):
        return val
    try:
        return datetime.strptime(str(val), "%Y-%m-%d").date()
    except Exception:
        return None

def doc_id_from_path(file_path):
    """Extract doc_id from filename stem, e.g. '/tier2/DOC-016.pdf' → 'DOC-016'."""
    basename = os.path.basename(file_path)
    return os.path.splitext(basename)[0]

def meta_defaults_from_lookup(doc_id, tier, file_path):
    """
    Return a metadata dict from the eds_synthetic.documents lookup,
    falling back to sensible defaults when not found.
    """
    row = METADATA_LOOKUP.get(doc_id)
    tier_num = safe_int(tier.replace("tier", ""), 4)
    if row:
        return {
            "doc_id":           doc_id,
            "title":            row["title"] or doc_id,
            "doc_type":         row["doc_type"] or "Document",
            "classification":   row["classification"] or "INTERNAL",
            "access_tier_level": row["access_tier_level"] or tier_num,
            "business_area":    row["business_area"] or "",
            "effective_date":   safe_date(row["effective_date"]),
            "author":           row["author"] or "",
            "version":          row["version"] or "v1.0",
            "source_system":    row["source_system"] or "Volume",
            "is_synthetic":     bool(row["is_synthetic"]) if row["is_synthetic"] is not None else True,
        }
    # Not in lookup — construct defaults from tier and path
    return {
        "doc_id":           doc_id,
        "title":            doc_id,
        "doc_type":         "Document",
        "classification":   {1: "RESTRICTED", 2: "CONFIDENTIAL", 3: "INTERNAL", 4: "PUBLIC"}.get(tier_num, "INTERNAL"),
        "access_tier_level": tier_num,
        "business_area":    "",
        "effective_date":   None,
        "author":           "",
        "version":          "v1.0",
        "source_system":    "Volume",
        "is_synthetic":     False,
    }


# ---- TXT parser (unchanged logic from previous version) ----

def parse_txt(tier, file_path):
    """
    Parse a plain-text file with YAML-style frontmatter header.
    Uses dbutils.fs.head() as per original implementation (handles dbfs:// paths).
    Returns (metadata_dict, content_str).
    """
    # dbutils.fs.head expects dbfs: or /Volumes/... paths; Unity Catalog volumes
    # are accessible directly via both the POSIX path and the dbfs:/ prefix form.
    raw = dbutils.fs.head(file_path, 1_000_000)
    meta_kv, body = {}, raw

    m = FRONTMATTER_RE.match(raw)
    if m:
        for line in m.group(1).splitlines():
            if ":" in line:
                key, _, val = line.partition(":")
                meta_kv[key.strip()] = val.strip()
        body = raw[m.end():]

    doc_id = meta_kv.get("doc_id") or doc_id_from_path(file_path)
    tier_num = safe_int(tier.replace("tier", ""), 4)

    meta = {
        "doc_id":           doc_id,
        "title":            meta_kv.get("title", ""),
        "doc_type":         meta_kv.get("doc_type", ""),
        "classification":   meta_kv.get("classification", "INTERNAL"),
        "access_tier_level": safe_int(meta_kv.get("access_tier_level"), tier_num),
        "business_area":    meta_kv.get("business_area", ""),
        "effective_date":   safe_date(meta_kv.get("effective_date")),
        "author":           meta_kv.get("author", ""),
        "version":          meta_kv.get("version", "v1.0"),
        "source_system":    meta_kv.get("source_system", "Volume"),
        "is_synthetic":     meta_kv.get("is_synthetic", "false").lower() == "true",
    }
    return meta, body.strip()


# ---- PDF parser ----

def parse_pdf(tier, file_path):
    """
    Extract text from a PDF using pdfplumber.
    For pages with no or minimal text (chart-heavy pages), inserts a
    placeholder note so downstream chunking is aware of visual content.
    Returns (metadata_dict, content_str).
    """
    import pdfplumber

    doc_id = doc_id_from_path(file_path)
    meta   = meta_defaults_from_lookup(doc_id, tier, file_path)

    page_texts = []
    with pdfplumber.open(file_path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            text = text.strip()
            if len(text) < 40:
                # Very little text — page is likely chart/image-dominated
                page_texts.append(f"[Page {page_num}] [Chart: visual content on this page]")
            else:
                page_texts.append(f"[Page {page_num}]\n{text}")

    content = "\n\n".join(page_texts)
    return meta, content


# ---- PPTX parser ----

def parse_pptx(tier, file_path):
    """
    Extract text from a PowerPoint file using python-pptx.
    For each slide, iterates all shapes and extracts:
      - Text shapes: all paragraphs joined
      - Chart shapes: chart title + series names and values as a text table
    Returns (metadata_dict, content_str).
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    doc_id = doc_id_from_path(file_path)
    meta   = meta_defaults_from_lookup(doc_id, tier, file_path)

    slide_texts = []

    with open(file_path, "rb") as f:
        prs = Presentation(f)

    for slide_num, slide in enumerate(prs.slides, start=1):
        parts = [f"[Slide {slide_num}]"]
        for shape in slide.shapes:
            # Text content
            if shape.has_text_frame:
                text = "\n".join(
                    p.text.strip()
                    for p in shape.text_frame.paragraphs
                    if p.text.strip()
                )
                if text:
                    parts.append(text)
            # Chart content
            if shape.has_chart:
                chart = shape.chart
                chart_title = ""
                try:
                    chart_title = chart.chart_title.text_frame.text.strip()
                except Exception:
                    chart_title = "Chart"
                chart_lines = [f"Chart: {chart_title}"]
                try:
                    for series in chart.series:
                        series_name = ""
                        try:
                            series_name = series.name or "Series"
                        except Exception:
                            series_name = "Series"
                        try:
                            values = [
                                str(round(v, 2)) if v is not None else "N/A"
                                for v in series.values
                            ]
                            chart_lines.append(f"  {series_name}: {', '.join(values)}")
                        except Exception:
                            chart_lines.append(f"  {series_name}: [values not accessible]")
                except Exception:
                    chart_lines.append("  [Series data not accessible]")
                parts.append("\n".join(chart_lines))

        slide_texts.append("\n".join(parts))

    content = "\n\n".join(slide_texts)
    return meta, content


# ---- XLSX parser ----

def parse_xlsx(tier, file_path):
    """
    Extract text from an Excel workbook using openpyxl.
    For each sheet:
      - Reads up to MAX_ROWS rows as a formatted text table (header + data rows)
      - For embedded chart objects, appends a descriptive note
    Returns (metadata_dict, content_str).
    """
    import openpyxl

    MAX_ROWS_PER_SHEET = 50
    doc_id = doc_id_from_path(file_path)
    meta   = meta_defaults_from_lookup(doc_id, tier, file_path)

    sheet_texts = []

    with open(file_path, "rb") as f:
        wb = openpyxl.load_workbook(f, data_only=True)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts = [f"[Sheet: {sheet_name}]"]

        # Extract cell data as table
        rows_data = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_ROWS_PER_SHEET:
                parts.append(f"[...truncated at {MAX_ROWS_PER_SHEET} rows]")
                break
            # Skip completely empty rows
            if all(cell is None for cell in row):
                continue
            row_str = "\t".join(
                str(cell) if cell is not None else ""
                for cell in row
            )
            rows_data.append(row_str)

        if rows_data:
            parts.append("\n".join(rows_data))

        # Detect and describe embedded charts
        try:
            for chart_obj in ws._charts:
                chart_title = ""
                try:
                    if chart_obj.title:
                        if hasattr(chart_obj.title, "tx") and chart_obj.title.tx:
                            # DrawingML chart title
                            try:
                                chart_title = "".join(
                                    r.t for r in chart_obj.title.tx.rich.p[0].r
                                    if hasattr(r, "t")
                                )
                            except Exception:
                                chart_title = str(chart_obj.title)
                        else:
                            chart_title = str(chart_obj.title)
                except Exception:
                    chart_title = "Untitled"
                chart_type = type(chart_obj).__name__
                parts.append(f"[Chart: {chart_title or 'Untitled'} — {chart_type}]")
        except Exception:
            pass  # _charts not available in some workbook states

        sheet_texts.append("\n".join(parts))

    content = "\n\n".join(sheet_texts)
    return meta, content

# COMMAND ----------

# MAGIC %md
# MAGIC ## Step 5: Process all files and build staged rows

# COMMAND ----------

run_id        = str(uuid.uuid4())
ingest_ts     = datetime.utcnow()
pipeline_name = "eds_document_ingestion_v3_multiformat"

rows   = []
errors = []

PARSER_MAP = {
    ".txt":  parse_txt,
    ".pdf":  parse_pdf,
    ".pptx": parse_pptx,
    ".xlsx": parse_xlsx,
}

for tier, file_path, ext in volume_files:
    try:
        parser = PARSER_MAP[ext]
        meta, content = parser(tier, file_path)

        doc_id     = meta["doc_id"]
        word_count = len(content.split()) if content else 0
        char_count = len(content)

        rows.append(Row(
            ingest_id         = str(uuid.uuid4()),
            doc_id            = doc_id,
            title             = meta.get("title") or "",
            doc_type          = meta.get("doc_type") or "",
            classification    = meta.get("classification") or "INTERNAL",
            access_tier_level = safe_int(meta.get("access_tier_level"), 4),
            business_area     = meta.get("business_area") or "",
            effective_date    = safe_date(meta.get("effective_date")),
            author            = meta.get("author") or "",
            version           = meta.get("version") or "v1.0",
            content           = content,
            source_file       = file_path,
            source_system     = meta.get("source_system") or "Volume",
            word_count        = word_count,
            char_count        = char_count,
            ingest_timestamp  = ingest_ts,
            ingest_pipeline   = pipeline_name,
            ingest_run_id     = run_id,
            is_synthetic      = bool(meta.get("is_synthetic", False)),
            ingestion_status  = "staged",
        ))

        print(
            f"[OK] {doc_id:10s}  {ext:6s}  "
            f"{meta.get('classification', '?'):12s}  "
            f"tier{meta.get('access_tier_level', '?')}  "
            f"{word_count:6,d} words"
        )

    except Exception as e:
        errors.append((file_path, str(e)))
        print(f"[ERROR] {file_path}: {e}")

print(f"\nParsed : {len(rows)} documents")
if errors:
    print(f"Errors : {len(errors)}")
    for path, err in errors:
        print(f"  {os.path.basename(path)}: {err}")

# COMMAND ----------

# MAGIC %md
# MAGIC ## Step 6: Write to staged_documents (upsert on doc_id)

# COMMAND ----------

schema = StructType([
    StructField("ingest_id",          StringType(),    False),
    StructField("doc_id",             StringType(),    False),
    StructField("title",              StringType(),    True),
    StructField("doc_type",           StringType(),    True),
    StructField("classification",     StringType(),    True),
    StructField("access_tier_level",  IntegerType(),   True),
    StructField("business_area",      StringType(),    True),
    StructField("effective_date",     DateType(),      True),
    StructField("author",             StringType(),    True),
    StructField("version",            StringType(),    True),
    StructField("content",            StringType(),    True),
    StructField("source_file",        StringType(),    True),
    StructField("source_system",      StringType(),    True),
    StructField("word_count",         IntegerType(),   True),
    StructField("char_count",         IntegerType(),   True),
    StructField("ingest_timestamp",   TimestampType(), True),
    StructField("ingest_pipeline",    StringType(),    True),
    StructField("ingest_run_id",      StringType(),    True),
    StructField("is_synthetic",       BooleanType(),   True),
    StructField("ingestion_status",   StringType(),    True),
])

if rows:
    staged_df = spark.createDataFrame(rows, schema=schema)
    staged_df.createOrReplaceTempView("staged_source")

    result = spark.sql(f"""
        MERGE INTO {CATALOG}.eds_raw.staged_documents AS target
        USING staged_source AS source
        ON target.doc_id = source.doc_id
        WHEN MATCHED THEN UPDATE SET
            target.ingest_id         = source.ingest_id,
            target.title             = source.title,
            target.doc_type          = source.doc_type,
            target.classification    = source.classification,
            target.access_tier_level = source.access_tier_level,
            target.business_area     = source.business_area,
            target.effective_date    = source.effective_date,
            target.author            = source.author,
            target.version           = source.version,
            target.content           = source.content,
            target.source_file       = source.source_file,
            target.source_system     = source.source_system,
            target.word_count        = source.word_count,
            target.char_count        = source.char_count,
            target.ingest_timestamp  = source.ingest_timestamp,
            target.ingest_pipeline   = source.ingest_pipeline,
            target.ingest_run_id     = source.ingest_run_id,
            target.is_synthetic      = source.is_synthetic,
            target.ingestion_status  = source.ingestion_status
        WHEN NOT MATCHED THEN INSERT *
    """)
    print(f"[OK] MERGE complete — {len(rows)} documents upserted into staged_documents")
else:
    print("[WARN] No rows to merge — check volume file discovery above")

# COMMAND ----------

# MAGIC %md
# MAGIC ## Step 7: Validation summary

# COMMAND ----------

staged_total = spark.table(f"{CATALOG}.eds_raw.staged_documents").count()

print("\n" + "=" * 65)
print("DOCUMENT INGESTION — Complete")
print("=" * 65)
print(f"  Source             : {VOLUME_ROOT}")
print(f"  Files discovered   : {len(volume_files)}")
print(f"  Documents parsed   : {len(rows)}")
print(f"  Errors             : {len(errors)}")
print(f"  staged_documents   : {staged_total} total rows")
print(f"  Pipeline run ID    : {run_id}")
print("=" * 65)

# COMMAND ----------

# MAGIC %md
# MAGIC ### Summary by format type

# COMMAND ----------

print("\nThis run — breakdown by file format:\n")
from collections import Counter, defaultdict

fmt_counter  = Counter()
tier_counter = Counter()
cls_counter  = Counter()

for row in rows:
    ext = os.path.splitext(row.source_file)[1].lower()
    fmt_counter[ext] += 1
    tier_counter[f"tier{row.access_tier_level}"] += 1
    cls_counter[row.classification] += 1

print("  By format:")
for fmt, cnt in sorted(fmt_counter.items()):
    print(f"    {fmt:8s}  {cnt} documents")

print("\n  By access tier:")
for tier_label, cnt in sorted(tier_counter.items()):
    print(f"    {tier_label}  {cnt} documents")

print("\n  By classification:")
for cls, cnt in sorted(cls_counter.items()):
    print(f"    {cls:14s}  {cnt} documents")

# COMMAND ----------

# MAGIC %md
# MAGIC ### Summary by classification and tier (from full table)

# COMMAND ----------

spark.sql(f"""
    SELECT
        classification,
        access_tier_level,
        COUNT(*)                         AS doc_count,
        ROUND(AVG(word_count), 0)        AS avg_words,
        ROUND(AVG(char_count) / 1000, 1) AS avg_chars_k,
        COUNT(DISTINCT ingest_pipeline)  AS pipelines
    FROM {CATALOG}.eds_raw.staged_documents
    GROUP BY classification, access_tier_level
    ORDER BY access_tier_level, classification
""").show()

# COMMAND ----------

# MAGIC %md
# MAGIC ### Per-document inventory (this run)

# COMMAND ----------

if rows:
    inv_rows = [(r.doc_id, r.classification, r.access_tier_level,
                 os.path.splitext(r.source_file)[1].lower(),
                 r.word_count, r.char_count) for r in rows]
    inv_df = spark.createDataFrame(
        inv_rows,
        schema=["doc_id", "classification", "access_tier_level", "format", "word_count", "char_count"]
    )
    inv_df.orderBy("access_tier_level", "doc_id").show(50, truncate=False)

print("\n[OK] Ingestion complete. Ready for 02_chunk_and_embed.")
