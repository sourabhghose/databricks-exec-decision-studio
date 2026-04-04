"""
POST /api/export/pptx  — Export a chat response as a PowerPoint file
"""

import io
import re
from datetime import datetime

import requests
from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

router = APIRouter()


class ExportRequest(BaseModel):
    content: str
    agent: str = "Strategic Intelligence"
    sources: list = []
    title: str = "Executive Decision Studio"


def _clean_for_pptx(text: str) -> str:
    """Strip markdown syntax for plain-text PPTX content."""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"^\|.+\|$", "", text, flags=re.MULTILINE)  # strip table rows
    text = re.sub(r"^\|[-:\s|]+\|$", "", text, flags=re.MULTILINE)  # strip separators
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


@router.post("/api/export/pptx")
async def export_pptx(req: ExportRequest):
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor

        ORANGE = RGBColor(0xF4, 0x79, 0x20)
        DARK   = RGBColor(0x1E, 0x29, 0x3B)
        GRAY   = RGBColor(0x64, 0x74, 0x8B)
        WHITE  = RGBColor(0xF1, 0xF5, 0xF9)

        prs = Presentation()
        prs.slide_width  = Inches(13.33)
        prs.slide_height = Inches(7.5)
        blank = prs.slide_layouts[6]  # blank layout

        def add_text(slide, text: str, left: float, top: float, w: float, h: float,
                     size: int, bold: bool = False, color: RGBColor = DARK, wrap: bool = True):
            txb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
            tf  = txb.text_frame
            tf.word_wrap = wrap
            for idx, line in enumerate(text.split("\n")):
                para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                run  = para.add_run()
                run.text = line
                run.font.size  = Pt(size)
                run.font.bold  = bold
                run.font.color.rgb = color

        # ── Title slide ───────────────────────────────────────────────────────
        s = prs.slides.add_slide(blank)
        add_text(s, "ALINTA ENERGY", 0.6, 0.6, 12, 0.6, 13, bold=True, color=ORANGE)
        add_text(s, "Executive Decision Studio", 0.6, 1.3, 12, 1.1, 30, bold=True, color=DARK)
        agent_nice = req.agent.replace("_", " ").title()
        add_text(s, agent_nice, 0.6, 2.7, 12, 0.6, 18, color=ORANGE)
        add_text(s, datetime.now().strftime("%d %B %Y") + "  ·  Confidential",
                 0.6, 3.4, 12, 0.5, 12, color=GRAY)

        # ── Content slides — split by ## headings ─────────────────────────────
        sections = re.split(r"\n(?=#{1,3} )", _clean_for_pptx(req.content))
        for section in sections:
            section = section.strip()
            if not section:
                continue
            lines = section.split("\n")
            heading  = re.sub(r"^#+\s*", "", lines[0]) if lines[0].startswith("#") else "Summary"
            body_raw = "\n".join(lines[1:] if lines[0].startswith("#") else lines).strip()
            if not heading and not body_raw:
                continue

            s = prs.slides.add_slide(blank)
            # Orange top bar
            bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.08))
            bar.fill.solid()
            bar.fill.fore_color.rgb = ORANGE
            bar.line.fill.background()

            add_text(s, heading, 0.4, 0.18, 12.5, 0.8, 22, bold=True, color=DARK)

            if body_raw:
                # Convert bullet lines
                body_lines = []
                for l in body_raw.split("\n"):
                    if l.startswith("- ") or l.startswith("* "):
                        body_lines.append("  •  " + l[2:])
                    else:
                        body_lines.append(l)
                add_text(s, "\n".join(body_lines), 0.4, 1.1, 12.5, 5.9, 13,
                         color=DARK, wrap=True)

        # ── Sources slide ─────────────────────────────────────────────────────
        if req.sources:
            s = prs.slides.add_slide(blank)
            bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.08))
            bar.fill.solid()
            bar.fill.fore_color.rgb = ORANGE
            bar.line.fill.background()
            add_text(s, "Source Documents", 0.4, 0.18, 12.5, 0.8, 22, bold=True, color=DARK)
            add_text(s, "\n".join(f"  •  {src}" for src in req.sources),
                     0.4, 1.1, 12.5, 5.9, 14, color=DARK)

        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)

        filename = f"executive-briefing-{datetime.now().strftime('%Y-%m-%d')}.pptx"
        return StreamingResponse(
            io.BytesIO(buf.read()),
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    except ImportError:
        return {"error": "python-pptx not installed. Add it to requirements.txt."}
    except Exception as e:
        print(f"[export] PPTX error: {e}")
        return {"error": str(e)}
